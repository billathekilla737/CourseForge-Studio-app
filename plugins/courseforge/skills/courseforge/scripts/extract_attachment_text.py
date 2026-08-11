"""extract_attachment_text.py (courseforge blind grading) - TEXT ONLY extraction.

Pulls plain text out of a submission attachment so the grading gateway can scrub
and pseudonymize it like a text-entry body. Deliberately narrow:

  .docx  paragraphs + table cells (python-docx)
  .pdf   per-page extracted text (pypdf); a scanned/image PDF yields little or
         nothing, which is reported rather than guessed at
  .txt   raw text

Everything else (images, pptx, xlsx, zip, code files, ...) is REFUSED - exit 3 -
so the caller lists it by filename only. Images especially stay out: screenshots
routinely contain names in window title bars, email headers, and file paths, and
no text scrubber sees pixels.

Output: extracted text on stdout (UTF-8), capped at MAX_CHARS with an explicit
truncation marker. Exit: 0 ok, 2 extraction produced no text, 3 unsupported type.
"""
import os
import sys

sys.stdout.reconfigure(encoding="utf-8", errors="replace")

MAX_CHARS = 120_000


def extract_docx(path):
    from docx import Document
    doc = Document(path)
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:
        for row in table.rows:
            parts.extend(cell.text for cell in row.cells)
    return "\n".join(p for p in parts if p and p.strip())


def extract_pdf(path):
    from pypdf import PdfReader
    reader = PdfReader(path)
    pages = []
    for pg in reader.pages:
        try:
            pages.append(pg.extract_text() or "")
        except Exception:
            pages.append("")
    return "\n".join(pages)


def extract_txt(path):
    with open(path, "r", encoding="utf-8", errors="replace") as f:
        return f.read()


def main():
    path = sys.argv[1]
    ext = os.path.splitext(path)[1].lower()
    handlers = {".docx": extract_docx, ".pdf": extract_pdf, ".txt": extract_txt}
    fn = handlers.get(ext)
    if fn is None:
        print("UNSUPPORTED:%s" % ext, file=sys.stderr)
        return 3
    try:
        text = fn(path)
    except Exception as e:
        print("EXTRACT FAILED: %s" % e, file=sys.stderr)
        return 2
    text = (text or "").strip()
    if not text:
        return 2
    if len(text) > MAX_CHARS:
        text = text[:MAX_CHARS] + "\n[TRUNCATED at %d chars]" % MAX_CHARS
    sys.stdout.write(text)
    return 0


if __name__ == "__main__":
    sys.exit(main())
